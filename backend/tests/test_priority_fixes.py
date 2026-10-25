"""Security boundaries, foreign-key deletion, review replacement and DOCX regressions."""
import json

import pytest
from docx import Document
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app.core.db import SessionLocal
from app.core.security import decode_token
from app.models import ContractFile, Message, ReviewResult, Session, SessionTypeEnum, User
from app.models.comparison import ComparisonTask
from app.models.review import ReviewTask


@pytest.fixture()
def foreign_keys(client):
    with SessionLocal() as db:
        db.execute(text("PRAGMA foreign_keys=ON"))
        db.commit()
    yield
    with SessionLocal() as db:
        db.execute(text("PRAGMA foreign_keys=OFF"))
        db.commit()


@pytest.fixture()
def other_headers(client):
    client.post("/api/users", json={"username": "other", "password": "secret123"})
    response = client.post("/api/auth/login", json={"identifier": "other", "password": "secret123"})
    return {"Authorization": "Bearer " + response.json()["data"]["access_token"]}


def user_id(headers):
    return decode_token(headers["Authorization"][7:])["user_id"]


@pytest.fixture()
def contract(auth_headers, tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("第一条 付款约定，应按合同履行交付与验收义务。")
    doc.save(path)
    parsed = tmp_path / "contract.txt"
    parsed.write_text("第一条 付款期限\n甲方在验收合格后支付合同价款。", encoding="utf-8")
    with SessionLocal() as db:
        record = ContractFile(
            user_id=user_id(auth_headers), title=path.name, file_type="docx",
            file_path=str(path), content_path=str(parsed), parse_status="parsed",
        )
        db.add(record)
        db.commit()
        return record.id, path, parsed


@pytest.mark.parametrize(("method", "path"), [
    ("get", "/api/users"), ("get", "/api/users/1"),
    ("put", "/api/users/1"), ("post", "/api/users/1/disable"),
    ("post", "/api/prompts"), ("get", "/api/model-configs"),
    ("post", "/api/contract-types"), ("get", "/api/dashboard/overview"),
    ("get", "/api/contracts/1/download"),
])
def test_anonymous_operations_are_rejected(client, method, path):
    assert getattr(client, method)(path).status_code == 401


@pytest.mark.parametrize(("method", "path"), [
    ("get", "/api/users"), ("post", "/api/users/1/disable"),
    ("post", "/api/contract-types"), ("put", "/api/contract-types/1"),
    ("post", "/api/contract-types/1/activate"), ("get", "/api/prompts"),
    ("post", "/api/model-configs"), ("delete", "/api/model-configs/1"),
])
def test_management_requires_admin(client, auth_headers, method, path):
    assert getattr(client, method)(path, headers=auth_headers).status_code == 403


def test_user_cannot_reset_or_read_other_account(client, auth_headers, other_headers):
    target = user_id(auth_headers)
    assert client.put(f"/api/users/{target}", headers=other_headers,
                      json={"password": "new-secret"}).status_code == 404
    assert client.get(f"/api/users/{target}", headers=other_headers).status_code == 404
    assert client.get(f"/api/users/{target}", headers=auth_headers).status_code == 200


def test_password_change_revokes_access_and_refresh(client, auth_headers):
    login = client.post("/api/auth/login", json={"identifier": "tester", "password": "secret123"}).json()["data"]
    headers = {"Authorization": "Bearer " + login["access_token"]}
    target = user_id(headers)
    response = client.put(f"/api/users/{target}", headers=headers, json={"password": "new-secret"})
    assert response.status_code == 200
    assert "x-access-token" not in response.headers
    assert client.get("/api/auth/me", headers=headers).status_code == 401
    assert client.post("/api/auth/refresh", json={"refresh_token": login["refresh_token"]}).status_code == 401
    assert client.post("/api/auth/login", json={"identifier": "tester", "password": "new-secret"}).status_code == 200


def test_contract_ownership_at_every_entry(client, auth_headers, other_headers, contract):
    file_id, _, _ = contract
    assert client.get(f"/api/contracts/{file_id}/download", headers=auth_headers).status_code == 200
    assert client.get(f"/api/contracts/{file_id}/download", headers=other_headers).status_code == 404
    for kind in ("chat", "review", "compare"):
        assert client.post("/api/sessions", headers=other_headers, json={
            "title": "foreign", "session_type": kind, "file_id": file_id,
        }).status_code == 404
    assert client.post("/api/comparisons", headers=other_headers, json={
        "standard_file_id": file_id, "comparison_file_id": file_id,
    }).status_code == 404
    assert client.post("/api/reviews/contracts/accept", headers=other_headers,
                       json={"file_id": file_id, "is_accepted": 1}).status_code == 404
    assert client.post("/api/reviews/contracts/accept", headers=auth_headers,
                       json={"file_id": file_id, "is_accepted": 1}).status_code == 200


def seed_review(db, owner, file_id):
    session = Session(user_id=owner, title="review", session_type=SessionTypeEnum.REVIEW, file_id=file_id)
    db.add(session)
    db.flush()
    task = ReviewTask(user_id=owner, session_id=session.id, file_id=file_id, status="completed")
    db.add(task)
    db.flush()
    db.add(ReviewResult(task_id=task.id, session_id=session.id, index=1, original_content="previous result"))
    db.commit()
    return session.id, task.id


def test_risk_acceptance_checks_task_owner(client, auth_headers, other_headers, contract):
    with SessionLocal() as db:
        sid, tid = seed_review(db, user_id(auth_headers), contract[0])
    body = {"session_id": sid, "task_id": tid, "index": 1, "is_accepted": 1}
    assert client.post("/api/reviews/risk-points/accept", headers=other_headers, json=body).status_code == 404
    assert client.post("/api/reviews/risk-points/accept", headers=auth_headers, json=body).status_code == 200


def test_legacy_foreign_session_cannot_read_contract(client, auth_headers, other_headers, contract):
    with SessionLocal() as db:
        ids = []
        for kind in (SessionTypeEnum.CHAT, SessionTypeEnum.REVIEW):
            session = Session(user_id=user_id(other_headers), title="legacy", session_type=kind, file_id=contract[0])
            db.add(session)
            db.flush()
            ids.append(session.id)
        db.commit()
    assert client.post("/api/chats", headers=other_headers, json={"session_id": ids[0], "content": "secret?"}).status_code == 404
    assert client.post("/api/reviews/start", headers=other_headers, json={"session_id": ids[1], "stance": "甲方"}).status_code == 404


def test_referenced_contract_is_preserved_and_session_cleanup_works(
    client, auth_headers, contract, foreign_keys,
):
    file_id, path, parsed = contract
    with SessionLocal() as db:
        sid, tid = seed_review(db, user_id(auth_headers), file_id)
        db.add(Message(session_id=sid, role="user", content="saved"))
        db.add(ComparisonTask(user_id=user_id(auth_headers), session_id=sid,
                              standard_file_id=file_id, comparison_file_id=file_id))
        db.commit()
    assert client.delete(f"/api/contracts/{file_id}", headers=auth_headers).status_code == 409
    assert path.exists() and parsed.exists()
    assert client.delete(f"/api/sessions/{sid}", headers=auth_headers).status_code == 200
    with SessionLocal() as db:
        assert db.get(ReviewTask, tid) is None
        assert db.query(Message).filter_by(session_id=sid).count() == 0
        assert db.query(ComparisonTask).filter_by(session_id=sid).count() == 0
    assert client.delete(f"/api/contracts/{file_id}", headers=auth_headers).status_code == 200
    assert not path.exists() and not parsed.exists()


def test_database_delete_failure_never_removes_files(client, auth_headers, contract, monkeypatch):
    from sqlalchemy.orm import Session as DBSession

    def fail(_):
        raise IntegrityError("delete", {}, Exception("foreign key"))

    monkeypatch.setattr(DBSession, "commit", fail)
    response = client.delete(f"/api/contracts/{contract[0]}", headers=auth_headers)
    assert response.status_code == 409
    assert contract[1].exists() and contract[2].exists()
    with SessionLocal() as db:
        assert db.get(ContractFile, contract[0]) is not None


async def test_review_notes_cannot_change_explicit_review_intent(monkeypatch):
    from app.agent.nodes import intent

    async def unexpected(_):
        pytest.fail("Explicit review requests must not be reclassified")

    monkeypatch.setattr(intent, "_llm_classify", unexpected)
    result = await intent.intent_node({"intent": "review", "description": "重点关注付款期限"})
    assert result["intent"] == "review"


@pytest.mark.parametrize("outcome", ["locked", "failure", "cancel", "expired", "empty", "commit_failure", "success"])
async def test_review_preserves_previous_until_success(
    client, auth_headers, contract, monkeypatch, foreign_keys, outcome,
):
    from app.schemas.review import ReviewStartRequest
    from app.services import review_service

    class Graph:
        async def astream(self, *args, **kwargs):
            if outcome == "failure":
                raise RuntimeError("unavailable")
            if outcome == "empty":
                return
            yield "updates", {"arbitration": {
                "final_risk_points": [{"original_content": "new result", "risk_level": "中"}],
                "summary": {"review_status": "completed"},
            }}

    released = []

    async def acquire(*args, **kwargs):
        return outcome != "locked"

    async def release(*args):
        released.append(args)

    async def renew(*args, **kwargs):
        return outcome != "expired"

    monkeypatch.setattr(review_service, "acquire_lock", acquire)
    monkeypatch.setattr(review_service, "release_lock", release)
    monkeypatch.setattr(review_service, "renew_lock", renew)
    monkeypatch.setattr(review_service, "get_graph", lambda _: Graph())
    with SessionLocal() as db:
        sid, tid = seed_review(db, user_id(auth_headers), contract[0])
        response = await review_service.stream_review_events(
            ReviewStartRequest(session_id=sid, stance="甲方"),
            db.get(User, user_id(auth_headers)), db, rag=object(),
        )
        if outcome == "commit_failure":
            commit = db.commit
            commits = 0

            def fail_result_commit():
                nonlocal commits
                commits += 1
                if commits == 2:
                    raise IntegrityError("insert", {}, Exception("database failure"))
                return commit()

            monkeypatch.setattr(db, "commit", fail_result_commit)
        assert db.query(ReviewTask).count() == 1
        if outcome == "cancel":
            await anext(response.body_iterator)
            await response.body_iterator.aclose()
        else:
            events = [json.loads(frame[6:]) async for frame in response.body_iterator]
            if outcome == "locked":
                assert events[0]["event"] == "error"
            if outcome == "success":
                assert events[-1]["event"] == "end"
        db.expire_all()
        points = db.query(ReviewResult).all()
        assert [p.original_content for p in points] == ["new result" if outcome == "success" else "previous result"]
        if outcome in ("failure", "cancel", "expired", "empty", "commit_failure"):
            latest = db.query(ReviewTask).order_by(ReviewTask.id.desc()).first()
            assert latest.status == ("cancelled" if outcome == "cancel" else "failed")
        assert len(released) == (0 if outcome == "locked" else 1)


def test_docx_tables_and_shifted_equal_indices(tmp_path):
    from app.api.comparisons import _diff_docs

    paths = [tmp_path / "a.docx", tmp_path / "b.docx"]
    for index, path in enumerate(paths):
        doc = Document()
        if index:
            doc.add_paragraph("新增说明")
        doc.add_paragraph("共同条款")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = f"金额 {index + 1} 万元"
        doc.add_paragraph("末尾条款")
        doc.save(path)
    summary, diffs = _diff_docs(*(str(p) for p in paths))
    assert summary.difference_count == 2
    changed = [d for d in diffs if d.operation == "replace"]
    assert len(changed) == 1
    assert "金额 1 万元" in changed[0].standard_text
    assert "金额 2 万元" in changed[0].comparison_text
    common = next(d for d in diffs if d.standard_text == "共同条款")
    assert common.std_index == 0 and common.cmp_index == 1


def test_table_only_changes_are_detected(tmp_path):
    from app.api.comparisons import _diff_docs

    paths = [tmp_path / "a.docx", tmp_path / "b.docx"]
    for index, path in enumerate(paths):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = str(index)
        doc.save(path)
    assert _diff_docs(*(str(p) for p in paths))[0].difference_count == 1


def test_admin_disable_revokes_session(client, admin_headers, other_headers):
    target = user_id(other_headers)
    assert client.post(f"/api/users/{target}/disable", headers=admin_headers).status_code == 200
    assert client.get("/api/auth/me", headers=other_headers).status_code == 401


def test_reused_comparison_session_updates_file_metadata(client, auth_headers, contract):
    file_id, path, _ = contract
    with SessionLocal() as db:
        replacement = ContractFile(user_id=user_id(auth_headers), title="replacement.docx",
                                   file_type="docx", file_path=str(path))
        db.add(replacement)
        db.commit()
        new_id = replacement.id
    first = client.post("/api/comparisons", headers=auth_headers, json={
        "standard_file_id": file_id, "comparison_file_id": file_id,
    }).json()["data"]
    second = client.post("/api/comparisons", headers=auth_headers, json={
        "standard_file_id": new_id, "comparison_file_id": file_id,
        "session_id": first["session_id"],
    })
    assert second.status_code == 200
    history = client.get(f"/api/sessions/{first['session_id']}/history", headers=auth_headers).json()["data"]
    assert history["standard_file"]["file_id"] == new_id
    assert history["standard_file"]["title"] == "replacement.docx"


async def test_nonreview_arbitration_is_incomplete(monkeypatch):
    from app.agent.nodes import arbitration

    async def summary(*args):
        return {"summary": "", "suggestion": "", "overall_risk": "低"}

    monkeypatch.setattr(arbitration, "_llm_summary", summary)
    result = await arbitration.arbitration_node({"gated": [], "stance": "甲方", "intent": "chat"})
    assert result["summary"]["review_status"] == "incomplete"


async def test_lock_renewal_cannot_extend_another_owner(client):
    from app.core.cache import acquire_lock, release_lock, renew_lock
    from app.core.redis import get_redis

    assert await acquire_lock("audit-lock", "new-owner", ttl=60)
    assert not await renew_lock("audit-lock", "expired-owner", ttl=900)
    await release_lock("audit-lock", "expired-owner")
    assert await get_redis().get("audit-lock") == "new-owner"
    assert await renew_lock("audit-lock", "new-owner", ttl=900)
    await release_lock("audit-lock", "new-owner")


async def test_second_review_cannot_mutate_first_active_attempt(
    client, auth_headers, contract, monkeypatch, foreign_keys,
):
    from app.schemas.review import ReviewStartRequest
    from app.services import review_service

    with SessionLocal() as first_db, SessionLocal() as second_db:
        sid, tid = seed_review(first_db, user_id(auth_headers), contract[0])
        request = ReviewStartRequest(session_id=sid, stance="甲方")
        first = await review_service.stream_review_events(
            request, first_db.get(User, user_id(auth_headers)), first_db, rag=object(),
        )
        first_event = json.loads((await anext(first.body_iterator))[6:])
        assert first_event["event"] == "meta"
        second = await review_service.stream_review_events(
            request, second_db.get(User, user_id(auth_headers)), second_db, rag=object(),
        )
        events = [json.loads(frame[6:]) async for frame in second.body_iterator]
        assert events[0]["event"] == "error"
        assert second_db.query(ReviewTask).count() == 2
        assert second_db.get(ReviewTask, first_event["data"]["task_id"]).status == "processing"
        assert second_db.query(ReviewResult).filter_by(task_id=tid).count() == 1
        await first.body_iterator.aclose()
